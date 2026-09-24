from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from models import TailoredApplication
from tailoring.ats_rules import FONT_NAME, FONT_SIZE_PT, HEADING_FONT_SIZE_PT
from tailoring.profile import Profile


def _set_base_fonts(doc: Document) -> None:
    """Single-column Word styles only — no tables/text-boxes are ever added,
    and every heading uses a real paragraph style (not manual bold+size) so
    ATS parsers that read style/outline info see genuine section breaks.
    """
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(FONT_SIZE_PT)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(HEADING_FONT_SIZE_PT)
        style.font.bold = True


def export_cv(profile: Profile, tailored: TailoredApplication, output_dir: Path) -> Path:
    doc = Document()
    _set_base_fonts(doc)

    contact = profile.contact
    doc.add_heading(contact.get("full_name", ""), level=1)
    contact_line = " | ".join(
        v for v in (contact.get("email"), contact.get("phone"), contact.get("location")) if v
    )
    if contact_line:
        doc.add_paragraph(contact_line)

    doc.add_heading("Personal Statement", level=2)
    doc.add_paragraph(tailored.personal_statement)

    doc.add_heading("Employment History", level=2)
    for entry in tailored.tailored_work_history:
        date_range = f"{entry.get('start_date', '')} – {entry.get('end_date', '')}"
        header = f"{entry.get('title', '')}, {entry.get('employer', '')} ({date_range})"
        doc.add_heading(header, level=3)
        location = entry.get("location", "")
        if location:
            doc.add_paragraph(location)
        for bullet in entry.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")

    if profile.education:
        doc.add_heading("Education & Qualifications", level=2)
        for edu in profile.education:
            date_range = f"{edu.get('start_date', '')} – {edu.get('end_date', '')}"
            line = f"{edu.get('qualification', '')}, {edu.get('institution', '')} ({date_range})"
            doc.add_paragraph(line, style="List Bullet")

    if profile.professional_registrations:
        doc.add_heading("Professional Registration", level=2)
        for reg in profile.professional_registrations:
            line = f"{reg.get('body', '')}: {reg.get('pin_or_number', '')}"
            if reg.get("expiry"):
                line += f" (expires {reg['expiry']})"
            doc.add_paragraph(line, style="List Bullet")

    if profile.skills:
        doc.add_heading("Skills", level=2)
        for skill in profile.skills:
            doc.add_paragraph(skill, style="List Bullet")

    if profile.training_cpd:
        doc.add_heading("Training & CPD", level=2)
        for item in profile.training_cpd:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("References", level=2)
    doc.add_paragraph("Available on request.")

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "CV.docx"
    doc.save(path)
    return path


def export_cover_letter(tailored: TailoredApplication, output_dir: Path) -> Path:
    doc = Document()
    _set_base_fonts(doc)

    for paragraph in tailored.cover_letter_or_supporting_statement.split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "cover_letter.docx"
    doc.save(path)
    return path


def export_supporting_statement_txt(tailored: TailoredApplication, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "supporting_statement.txt"
    path.write_text(tailored.cover_letter_or_supporting_statement, encoding="utf-8")
    return path


def export_criteria_responses_txt(tailored: TailoredApplication, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "criteria_responses.txt"
    lines = []
    for resp in tailored.criteria_responses:
        tag = "Essential" if resp.essential else "Desirable"
        lines.append(f"[{tag}] {resp.criterion}")
        lines.append(resp.evidence or "(no evidence found in profile — flagged as a gap)")
        lines.append("")
    if tailored.flagged_gaps:
        lines.append("--- Flagged gaps (no genuine evidence in your profile) ---")
        lines.extend(tailored.flagged_gaps)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def export_all(profile: Profile, tailored: TailoredApplication, output_dir: Path) -> dict[str, Path]:
    return {
        "cv": export_cv(profile, tailored, output_dir),
        "cover_letter": export_cover_letter(tailored, output_dir),
        "supporting_statement_txt": export_supporting_statement_txt(tailored, output_dir),
        "criteria_responses_txt": export_criteria_responses_txt(tailored, output_dir),
    }
